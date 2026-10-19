#!/usr/bin/env python3
"""
Atualiza articles.json com a categorização MANUAL da Célia (planilha nova).
Usa sheet "Nova Categorias acrescentadas" que tem os dados alinhados corretamente.
"""

import json
import re
from pathlib import Path
from openpyxl import load_workbook

EXCEL = Path("/home/z/my-project/upload/Nova Lista de Categorização.xlsx")
JSON_PATH = Path("/home/z/my-project/src/data/articles.json")

NEW_CONTEXTS = {
    "Dinâmica Familiar e Criação de Filhos": {
        "slug": "dinamica-familiar",
        "label": "Dinâmica Familiar e Criação de Filhos",
        "description": "Mães, pais, filhos, sogras, avós. A família como primeiro espelho e território de aprendizado.",
    },
    "Relacionamentos Amorosos e Sexualidade": {
        "slug": "relacionamentos-sexualidade",
        "label": "Relacionamentos Amorosos e Sexualidade",
        "description": "Amor, parceria, sexo, escolhas, separação. Tudo que envolve escolher e ser escolhido por outro.",
    },
    "Ansiedade, Estresse e Esgotamento": {
        "slug": "ansiedade-estresse",
        "label": "Ansiedade, Estresse e Esgotamento",
        "description": "Quando o corpo pede calma e a cabeça não consegue parar. Cansaço, exaustão, dificuldade de relaxar.",
    },
    "Autocrítica, Limites e Autoestima": {
        "slug": "autocritica-limites",
        "label": "Autocrítica, Limites e Autoestima",
        "description": "Aprender a dizer não, reconhecer o próprio valor, calar o juiz interior, cuidar de si sem culpa.",
    },
    "Sintomas Psicossomáticos e Saúde Emocional": {
        "slug": "sintomas-psicossomaticos-saude-emocional",
        "label": "Sintomas Psicossomáticos e Saúde Emocional",
        "description": "Emoção e corpo: como o que sentimos vira sintoma, doença como mensagem, terapias integrativas.",
    },
    "Produtividade, Carreira e Organização": {
        "slug": "produtividade-carreira",
        "label": "Produtividade, Carreira e Organização",
        "description": "Vocação, propósito, dinheiro como energia, trabalho como construção de sentido.",
    },
    "Espiritualidade e Autoconhecimento": {
        "slug": "espiritualidade-autoconhecimento",
        "label": "Espiritualidade e Autoconhecimento",
        "description": "Espiritualidade no cotidiano, meditação, exercícios de proteção, autoavaliação, paz interior.",
    },
}

# Manual fixes (artigos que não aparecem na planilha)
MANUAL_FIXES = {
    "o que é masturbação: entenda da puberdade à maturidade": {
        "contexto": "Relacionamentos Amorosos e Sexualidade",
        "tipo": "Reflexiva",
        "tags": "Sexualidade, Puberdade, Masturbação, Tabu Sexual",
    },
}


def normalize_title(t: str) -> str:
    t = t.lower().strip()
    t = re.sub(r'["\']', '', t)
    t = re.sub(r'^#+', '', t)
    t = re.sub(r'[:\?]+$', '', t)
    t = re.sub(r'\s+', ' ', t)
    return t.strip()


def main():
    wb = load_workbook(EXCEL)
    ws = wb["Nova Categorias acrescentadas"]
    manual = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        titulo = str(row[0]).strip()
        contexto = str(row[1]).strip() if row[1] else ""
        tipo = str(row[2]).strip() if row[2] else ""
        tags = str(row[3]).strip() if row[3] else ""
        manual[normalize_title(titulo)] = {
            "titulo_original": titulo,
            "contexto": contexto,
            "tipo": tipo,
            "tags": tags,
        }
    print(f"[1] Manual: {len(manual)} artigos catalogados")

    articles = json.loads(JSON_PATH.read_text(encoding='utf-8'))
    print(f"[2] JSON: {len(articles)} artigos")

    matched = 0
    unmatched = []
    for a in articles:
        key = normalize_title(a["title"])
        m = manual.get(key) or MANUAL_FIXES.get(key)
        if m:
            a["contentType"] = "pratica" if m["tipo"].lower().startswith("prát") else "artigo"
            ctx = m["contexto"]
            if ctx in NEW_CONTEXTS:
                ctx_data = NEW_CONTEXTS[ctx]
                a["contexts"] = [ctx_data["slug"]]
                a["contextLabels"] = [ctx_data["label"]]
                a["contextDescription"] = ctx_data["description"]
            else:
                a["contexts"] = ["outros"]
                a["contextLabels"] = [ctx or "Outros"]
            tag_list = [t.strip() for t in re.split(r'[,;]', m["tags"]) if t.strip()]
            a["tags"] = tag_list
            old_kw = a.get("keywords", []) if a.get("keywords") else []
            a["keywords"] = list(dict.fromkeys(old_kw + tag_list))
            a["contextOriginal"] = ctx
            a["tipoOriginal"] = m["tipo"]
            matched += 1
        else:
            unmatched.append(a["title"])

    print(f"[3] Match: {matched}/{len(articles)}")
    if unmatched:
        print(f"   ⚠ {len(unmatched)} sem match:")
        for t in unmatched[:3]:
            print(f"     - {t}")

    JSON_PATH.write_text(json.dumps(articles, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f"[4] JSON atualizado: {JSON_PATH}")

    from collections import defaultdict
    by_ctx = defaultdict(int)
    by_type = defaultdict(int)
    for a in articles:
        for c in a["contexts"]:
            by_ctx[c] += 1
        by_type[a["contentType"]] += 1
    print("\nDistribuição por tipo:")
    for t, n in sorted(by_type.items()):
        print(f"  {t}: {n}")
    print("\nDistribuição por contexto:")
    for c, n in sorted(by_ctx.items(), key=lambda x: -x[1]):
        print(f"  {c}: {n}")


if __name__ == "__main__":
    main()
